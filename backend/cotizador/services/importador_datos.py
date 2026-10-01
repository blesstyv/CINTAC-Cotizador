import sqlite3

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.db import transaction

from openpyxl import load_workbook

from cotizador.models import (
    Puerto,
    Ruta,
    Tarifa,
    TiempoTransito,
    TipoContenedor,
)


NOMBRE_HOJA = "Tarifas Referencia"

PAIS_DESTINO = "Chile"

TIPOS_RUTA_PERMITIDOS = {
    "Directo",
    "Transbordo",
    "Directo / Transbordo",
}

COLUMNAS_ESPERADAS = [
    "Clave",
    "Puerto Origen",
    "Pais Origen",
    "Puerto Destino",
    "Tipo de Ruta",
    "Tarifa 20' Min (US$)",
    "Tarifa 20' Max (US$)",
    "Tarifa 40' Min (US$)",
    "Tarifa 40' Max (US$)",
    "Transito Min (dias)",
    "Transito Max (dias)",
    "Fuente",
]


class ErrorImportacionExcel(Exception):
    def __init__(self, errores):
        self.errores = errores

        super().__init__(
            "El archivo contiene datos inválidos."
        )


@dataclass(frozen=True)
class FilaTarifa:
    numero_fila: int
    clave: str
    origen: str
    pais_origen: str
    destino: str
    tipo_ruta: str
    tarifa20_min: Decimal
    tarifa20_max: Decimal
    tarifa40_min: Decimal
    tarifa40_max: Decimal
    transito_min: int
    transito_max: int
    fuente: str


def _texto(valor):
    if valor is None:
        return ""

    return str(valor).strip()


def _validar_texto(
    valor,
    nombre_campo,
    numero_fila,
    errores,
):
    texto = _texto(
        valor
    )

    if not texto:
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' es obligatorio."
            )
        )

        return None

    if texto.startswith("="):
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' no puede "
                "contener una fórmula."
            )
        )

        return None

    return texto


def _convertir_decimal(
    valor,
    nombre_campo,
    numero_fila,
    errores,
    max_decimales=None,
    max_digitos=None,
):
    if (
        valor is None
        or valor == ""
    ):
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' es obligatorio."
            )
        )

        return None

    if isinstance(
        valor,
        bool,
    ):
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' debe ser numérico."
            )
        )

        return None

    try:
        if isinstance(
            valor,
            str,
        ):
            texto = (
                valor.strip()
            )

            if texto.startswith(
                "="
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        f"'{nombre_campo}' no puede "
                        "contener una fórmula."
                    )
                )

                return None

            if (
                "," in texto
                and "." in texto
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        f"'{nombre_campo}' posee "
                        "un formato numérico ambiguo."
                    )
                )

                return None

            texto = texto.replace(
                ",",
                ".",
            )

            numero = Decimal(
                texto
            )

        else:
            numero = Decimal(
                str(valor)
            )

    except (
        InvalidOperation,
        ValueError,
        TypeError,
    ):
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' debe "
                "contener un valor numérico."
            )
        )

        return None

    if not numero.is_finite():
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' debe "
                "contener un número finito."
            )
        )

        return None

    if numero <= 0:
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' debe "
                "ser mayor que cero."
            )
        )

        return None

    if (
        max_decimales
        is not None
    ):
        exponente = (
            numero
            .as_tuple()
            .exponent
        )

        decimales = max(
            -exponente,
            0,
        )

        if (
            decimales
            > max_decimales
        ):
            errores.append(
                (
                    f"Fila {numero_fila}: "
                    f"'{nombre_campo}' no puede "
                    f"tener más de "
                    f"{max_decimales} decimales."
                )
            )

            return None

    if (
        max_digitos
        is not None
    ):
        (
            _,
            digitos,
            exponente,
        ) = numero.as_tuple()

        cantidad_digitos = len(
            digitos
        )

        if exponente > 0:
            cantidad_digitos += (
                exponente
            )

        if (
            cantidad_digitos
            > max_digitos
        ):
            errores.append(
                (
                    f"Fila {numero_fila}: "
                    f"'{nombre_campo}' supera "
                    f"el máximo de "
                    f"{max_digitos} dígitos "
                    "permitidos."
                )
            )

            return None

    return numero


def _convertir_entero(
    valor,
    nombre_campo,
    numero_fila,
    errores,
):
    numero = (
        _convertir_decimal(
            valor,
            nombre_campo,
            numero_fila,
            errores,
        )
    )

    if numero is None:
        return None

    if (
        numero
        != numero.to_integral_value()
    ):
        errores.append(
            (
                f"Fila {numero_fila}: "
                f"'{nombre_campo}' debe "
                "ser un número entero."
            )
        )

        return None

    return int(
        numero
    )


def _fila_vacia(
    valores,
):
    return all(
        valor is None
        or (
            isinstance(
                valor,
                str,
            )
            and not valor.strip()
        )
        for valor
        in valores
    )


def validar_archivo_excel(
    ruta_archivo,
):
    ruta = Path(
        ruta_archivo
    )

    if not ruta.exists():
        raise ErrorImportacionExcel(
            [
                (
                    "El archivo indicado "
                    "no existe."
                )
            ]
        )

    if not ruta.is_file():
        raise ErrorImportacionExcel(
            [
                (
                    "La ruta indicada no "
                    "corresponde a un archivo."
                )
            ]
        )

    if (
        ruta.suffix.lower()
        != ".xlsx"
    ):
        raise ErrorImportacionExcel(
            [
                (
                    "El archivo debe tener "
                    "formato .xlsx."
                )
            ]
        )

    try:
        libro = load_workbook(
            filename=ruta,
            read_only=True,
            data_only=False,
        )

    except Exception as error:
        raise ErrorImportacionExcel(
            [
                (
                    "No fue posible abrir "
                    "el archivo Excel."
                )
            ]
        ) from error

    try:
        if (
            NOMBRE_HOJA
            not in libro.sheetnames
        ):
            raise ErrorImportacionExcel(
                [
                    (
                        "No existe la hoja "
                        f"'{NOMBRE_HOJA}'."
                    )
                ]
            )

        hoja = libro[
            NOMBRE_HOJA
        ]

        encabezados = [
            _texto(
                celda.value
            )
            for celda
            in hoja[1]
        ]

        if (
            encabezados
            != COLUMNAS_ESPERADAS
        ):
            errores = [
                (
                    "La estructura de columnas "
                    "no coincide con el formato "
                    "esperado de CINTAC."
                )
            ]

            faltantes = [
                columna
                for columna
                in COLUMNAS_ESPERADAS
                if columna
                not in encabezados
            ]

            adicionales = [
                columna
                for columna
                in encabezados
                if columna
                and columna
                not in COLUMNAS_ESPERADAS
            ]

            if faltantes:
                errores.append(
                    (
                        "Columnas faltantes: "
                        + ", ".join(
                            faltantes
                        )
                    )
                )

            if adicionales:
                errores.append(
                    (
                        "Columnas no reconocidas: "
                        + ", ".join(
                            adicionales
                        )
                    )
                )

            raise ErrorImportacionExcel(
                errores
            )

        errores = []
        filas_validas = []

        claves_vistas = {}
        rutas_vistas = {}
        pais_por_origen = {}

        for (
            numero_fila,
            fila,
        ) in enumerate(
            hoja.iter_rows(
                min_row=2,
                max_col=len(
                    COLUMNAS_ESPERADAS
                ),
                values_only=True,
            ),
            start=2,
        ):
            if _fila_vacia(
                fila
            ):
                continue

            errores_antes = len(
                errores
            )

            (
                valor_clave,
                valor_origen,
                valor_pais,
                valor_destino,
                valor_tipo_ruta,
                valor_tarifa20_min,
                valor_tarifa20_max,
                valor_tarifa40_min,
                valor_tarifa40_max,
                valor_transito_min,
                valor_transito_max,
                valor_fuente,
            ) = fila

            clave = (
                _validar_texto(
                    valor_clave,
                    "Clave",
                    numero_fila,
                    errores,
                )
            )

            origen = (
                _validar_texto(
                    valor_origen,
                    "Puerto Origen",
                    numero_fila,
                    errores,
                )
            )

            pais = (
                _validar_texto(
                    valor_pais,
                    "Pais Origen",
                    numero_fila,
                    errores,
                )
            )

            destino = (
                _validar_texto(
                    valor_destino,
                    "Puerto Destino",
                    numero_fila,
                    errores,
                )
            )

            tipo_ruta = (
                _validar_texto(
                    valor_tipo_ruta,
                    "Tipo de Ruta",
                    numero_fila,
                    errores,
                )
            )

            fuente = (
                _validar_texto(
                    valor_fuente,
                    "Fuente",
                    numero_fila,
                    errores,
                )
            )

            tarifa20_min = (
                _convertir_decimal(
                    valor_tarifa20_min,
                    "Tarifa 20' Min (US$)",
                    numero_fila,
                    errores,
                    max_decimales=2,
                    max_digitos=12,
                )
            )

            tarifa20_max = (
                _convertir_decimal(
                    valor_tarifa20_max,
                    "Tarifa 20' Max (US$)",
                    numero_fila,
                    errores,
                    max_decimales=2,
                    max_digitos=12,
                )
            )

            tarifa40_min = (
                _convertir_decimal(
                    valor_tarifa40_min,
                    "Tarifa 40' Min (US$)",
                    numero_fila,
                    errores,
                    max_decimales=2,
                    max_digitos=12,
                )
            )

            tarifa40_max = (
                _convertir_decimal(
                    valor_tarifa40_max,
                    "Tarifa 40' Max (US$)",
                    numero_fila,
                    errores,
                    max_decimales=2,
                    max_digitos=12,
                )
            )

            transito_min = (
                _convertir_entero(
                    valor_transito_min,
                    "Transito Min (dias)",
                    numero_fila,
                    errores,
                )
            )

            transito_max = (
                _convertir_entero(
                    valor_transito_max,
                    "Transito Max (dias)",
                    numero_fila,
                    errores,
                )
            )

            if (
                tipo_ruta
                and tipo_ruta
                not in TIPOS_RUTA_PERMITIDOS
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        "'Tipo de Ruta' posee "
                        "un valor no permitido: "
                        f"'{tipo_ruta}'."
                    )
                )

            if (
                origen
                and destino
                and origen.casefold()
                == destino.casefold()
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        "el puerto de origen "
                        "y destino no pueden "
                        "ser iguales."
                    )
                )

            if (
                clave
                and origen
                and destino
            ):
                clave_esperada = (
                    f"{origen}|{destino}"
                )

                if (
                    clave
                    != clave_esperada
                ):
                    errores.append(
                        (
                            f"Fila {numero_fila}: "
                            f"la clave '{clave}' "
                            "no coincide con "
                            "Puerto Origen + "
                            "Puerto Destino. "
                            "Debería ser "
                            f"'{clave_esperada}'."
                        )
                    )

            if clave:
                if (
                    clave
                    in claves_vistas
                ):
                    errores.append(
                        (
                            f"Fila {numero_fila}: "
                            f"la clave '{clave}' "
                            "está duplicada. "
                            "Ya apareció en la "
                            "fila "
                            f"{claves_vistas[clave]}."
                        )
                    )

                else:
                    claves_vistas[
                        clave
                    ] = numero_fila

            if (
                origen
                and destino
            ):
                identificador_ruta = (
                    origen.casefold(),
                    destino.casefold(),
                )

                if (
                    identificador_ruta
                    in rutas_vistas
                ):
                    errores.append(
                        (
                            f"Fila {numero_fila}: "
                            "la combinación "
                            f"'{origen} → "
                            f"{destino}' "
                            "está duplicada. "
                            "Ya apareció en la "
                            "fila "
                            f"{rutas_vistas[identificador_ruta]}."
                        )
                    )

                else:
                    rutas_vistas[
                        identificador_ruta
                    ] = numero_fila

            if (
                origen
                and pais
            ):
                clave_origen = (
                    origen.casefold()
                )

                if (
                    clave_origen
                    in pais_por_origen
                ):
                    pais_anterior = (
                        pais_por_origen[
                            clave_origen
                        ][0]
                    )

                    fila_anterior = (
                        pais_por_origen[
                            clave_origen
                        ][1]
                    )

                    if (
                        pais_anterior.casefold()
                        != pais.casefold()
                    ):
                        errores.append(
                            (
                                f"Fila {numero_fila}: "
                                f"el puerto '{origen}' "
                                "aparece asociado al "
                                f"país '{pais}', pero "
                                "anteriormente fue "
                                "asociado al país "
                                f"'{pais_anterior}' "
                                "en la fila "
                                f"{fila_anterior}."
                            )
                        )

                else:
                    pais_por_origen[
                        clave_origen
                    ] = (
                        pais,
                        numero_fila,
                    )

            if (
                tarifa20_min
                is not None
                and tarifa20_max
                is not None
                and tarifa20_max
                < tarifa20_min
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        "la tarifa máxima de "
                        "20' no puede ser menor "
                        "que la tarifa mínima."
                    )
                )

            if (
                tarifa40_min
                is not None
                and tarifa40_max
                is not None
                and tarifa40_max
                < tarifa40_min
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        "la tarifa máxima de "
                        "40' no puede ser menor "
                        "que la tarifa mínima."
                    )
                )

            if (
                transito_min
                is not None
                and transito_max
                is not None
                and transito_max
                < transito_min
            ):
                errores.append(
                    (
                        f"Fila {numero_fila}: "
                        "el tránsito máximo no "
                        "puede ser menor que "
                        "el tránsito mínimo."
                    )
                )

            if (
                len(
                    errores
                )
                == errores_antes
            ):
                filas_validas.append(
                    FilaTarifa(
                        numero_fila=
                            numero_fila,

                        clave=
                            clave,

                        origen=
                            origen,

                        pais_origen=
                            pais,

                        destino=
                            destino,

                        tipo_ruta=
                            tipo_ruta,

                        tarifa20_min=
                            tarifa20_min,

                        tarifa20_max=
                            tarifa20_max,

                        tarifa40_min=
                            tarifa40_min,

                        tarifa40_max=
                            tarifa40_max,

                        transito_min=
                            transito_min,

                        transito_max=
                            transito_max,

                        fuente=
                            fuente,
                    )
                )

        if errores:
            raise ErrorImportacionExcel(
                errores
            )

        if not filas_validas:
            raise ErrorImportacionExcel(
                [
                    (
                        "La hoja no contiene "
                        "filas de datos válidas."
                    )
                ]
            )

        return filas_validas

    finally:
        libro.close()


def crear_respaldo_sqlite():
    ruta_bd = Path(
        settings.DATABASES[
            "default"
        ]["NAME"]
    )

    if not ruta_bd.exists():
        return None

    carpeta_respaldos = (
        Path(
            settings.BASE_DIR
        )
        / "backups"
    )

    carpeta_respaldos.mkdir(
        parents=True,
        exist_ok=True,
    )

    marca_tiempo = (
        datetime.now()
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    ruta_respaldo = (
        carpeta_respaldos
        / (
            "db_antes_importacion_"
            f"{marca_tiempo}.sqlite3"
        )
    )

    origen = sqlite3.connect(
        str(
            ruta_bd
        )
    )

    destino = sqlite3.connect(
        str(
            ruta_respaldo
        )
    )

    try:
        origen.backup(
            destino
        )

    finally:
        destino.close()
        origen.close()

    return ruta_respaldo


@transaction.atomic
def aplicar_importacion(
    filas,
):
    resumen = {
        "puertos_creados": 0,
        "rutas_creadas": 0,
        "rutas_actualizadas": 0,
        "tarifas_creadas": 0,
        "tarifas_actualizadas": 0,
        "transitos_creados": 0,
        "transitos_actualizados": 0,
    }

    TipoContenedor.objects.get_or_create(
        codigo="20",
        defaults={
            "nombre":
                "Contenedor 20'",

            "capacidad_tn":
                Decimal("25.00"),

            "activo":
                True,
        },
    )

    TipoContenedor.objects.get_or_create(
        codigo="40",
        defaults={
            "nombre":
                "Contenedor 40'",

            "capacidad_tn":
                Decimal("25.00"),

            "activo":
                True,
        },
    )

    for fila in filas:
        origen, creado = (
            Puerto.objects.get_or_create(
                nombre=
                    fila.origen,

                pais=
                    fila.pais_origen,
            )
        )

        if creado:
            resumen[
                "puertos_creados"
            ] += 1

        destino, creado = (
            Puerto.objects.get_or_create(
                nombre=
                    fila.destino,

                pais=
                    PAIS_DESTINO,
            )
        )

        if creado:
            resumen[
                "puertos_creados"
            ] += 1

        ruta, creada = (
            Ruta.objects.update_or_create(
                puerto_origen=
                    origen,

                puerto_destino=
                    destino,

                defaults={
                    "tipo_ruta":
                        fila.tipo_ruta,
                },
            )
        )

        if creada:
            resumen[
                "rutas_creadas"
            ] += 1

        else:
            resumen[
                "rutas_actualizadas"
            ] += 1

        for (
            codigo,
            minimo,
            maximo,
        ) in [
            (
                "20",
                fila.tarifa20_min,
                fila.tarifa20_max,
            ),
            (
                "40",
                fila.tarifa40_min,
                fila.tarifa40_max,
            ),
        ]:
            _, creada = (
                Tarifa.objects.update_or_create(
                    ruta=
                        ruta,

                    tipo_contenedor=
                        codigo,

                    defaults={
                        "valor_minimo":
                            minimo,

                        "valor_maximo":
                            maximo,

                        "fuente":
                            fila.fuente,
                    },
                )
            )

            if creada:
                resumen[
                    "tarifas_creadas"
                ] += 1

            else:
                resumen[
                    "tarifas_actualizadas"
                ] += 1

        _, creado = (
            TiempoTransito.objects
            .update_or_create(
                ruta=
                    ruta,

                defaults={
                    "dias_minimos":
                        fila.transito_min,

                    "dias_maximos":
                        fila.transito_max,
                },
            )
        )

        if creado:
            resumen[
                "transitos_creados"
            ] += 1

        else:
            resumen[
                "transitos_actualizados"
            ] += 1

    return resumen